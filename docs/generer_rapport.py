"""Genere le rapport du projet au format Word, dans un style sobre et homogene.

Principes de mise en forme :
  - texte noir, une seule police, pas de couleur ;
  - phrases courtes, tableaux simples a bordures fines ;
  - les chiffres de performance sont lus dans le modele entraine, jamais
    recopies a la main (meme exigence que le critere CA-4 du cahier des charges).

Usage : python -m docs.generer_rapport
Produit : RAPPORT_AgroPredict_SN.docx
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from models.predict import get_model_info, get_model_metrics

NOIR = RGBColor(0x1A, 0x1A, 0x1A)
POLICE = "Calibri"
SORTIE = "RAPPORT_AgroPredict_SN.docx"


def _style_base(doc):
    """Police et couleur uniformes pour tout le document."""
    normal = doc.styles["Normal"]
    normal.font.name = POLICE
    normal.font.size = Pt(11)
    normal.font.color.rgb = NOIR
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for niveau, taille in [("Heading 1", 16), ("Heading 2", 13)]:
        s = doc.styles[niveau]
        s.font.name = POLICE
        s.font.size = Pt(taille)
        s.font.bold = True
        s.font.color.rgb = NOIR


def titre(doc, texte, niveau=1):
    doc.add_heading(texte, level=niveau)


def par(doc, texte, gras=False, italique=False, centre=False):
    p = doc.add_paragraph()
    r = p.add_run(texte)
    r.bold = gras
    r.italic = italique
    if centre:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def puces(doc, elements):
    for e in elements:
        doc.add_paragraph(e, style="List Bullet")


def tableau(doc, entetes, lignes):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    for i, e in enumerate(entetes):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].add_run(e).bold = True
    for ligne in lignes:
        cells = t.add_row().cells
        for i, valeur in enumerate(ligne):
            cells[i].text = str(valeur)
    doc.add_paragraph()
    return t


def construire():
    metriques = get_model_metrics()
    infos = get_model_info()
    if not metriques:
        raise SystemExit("Modele absent. Lancez d'abord : python -m models.train_model")

    doc = Document()
    _style_base(doc)

    # ---- Page de titre ----
    if os.path.exists("docs/logo.png"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture("docs/logo.png", width=Pt(260))

    doc.add_paragraph()
    par(doc, "Rapport de projet", gras=True, centre=True).runs[0].font.size = Pt(20)
    par(doc, "AgroPredict SN — Simulateur de rendement agricole", centre=True)
    doc.add_paragraph()
    par(doc, "Papa Malick NDIAYE", centre=True)
    par(doc, "Master 2 Data Science & Génie Logiciel", centre=True)
    par(doc, "Université Alioune Diop de Bambey", centre=True)
    par(doc, "njaymika@gmail.com", centre=True)
    par(doc, "Juillet 2026", centre=True)
    doc.add_page_break()

    # ---- 1. Présentation en bref ----
    titre(doc, "1. Présentation en bref")
    par(doc, "En une phrase", gras=True)
    par(doc, "Un agriculteur répond à cinq questions sur son téléphone, et "
             "l'application lui dit combien il va récolter.")
    par(doc, "Comment ça marche, en trois temps", gras=True)
    par(doc, "1. Les données. On part des rendements agricoles réels du Sénégal "
             "(FAOSTAT) et de la météo réelle (NASA POWER). Un script assemble tout "
             "cela en un tableau : région, culture, engrais, irrigation, et le "
             "rendement obtenu.")
    par(doc, "2. L'apprentissage. Un modèle XGBoost lit ce tableau et apprend les "
             "régularités : le riz de Saint-Louis rend plus que le mil de Sédhiou, "
             "l'engrais aide un peu. On enregistre ce qu'il a appris dans un fichier.")
    par(doc, "3. L'application. L'agriculteur choisit sa région, sa culture, sa "
             "surface, son engrais et son irrigation. Le modèle calcule, et l'écran "
             "affiche deux chiffres : le rendement en tonnes par hectare et la "
             "récolte totale en sacs.")
    par(doc, "Ce n'est pas un modèle qui prédit l'avenir. C'est un simulateur "
             "honnête, industrialisé proprement : données assumées, modèle traçable, "
             "tests qui protègent réellement.", italique=True)

    # ---- 2. Contexte ----
    titre(doc, "2. Contexte et problème")
    par(doc, "L'agriculture fait vivre une grande partie de la population "
             "sénégalaise. En début de saison, l'agriculteur décide quoi semer, sur "
             "quelle surface, avec ou sans engrais. Il prend cette décision sans "
             "repère chiffré.")
    par(doc, "AgroPredict SN répond à une seule question : pour ma culture, dans ma "
             "région, avec mes pratiques, quel rendement puis-je attendre ?")

    # ---- 3. Ce que fait l'application ----
    titre(doc, "3. Ce que l'application fait et ne fait pas")
    par(doc, "L'utilisateur renseigne cinq informations qu'il connaît :", gras=True)
    puces(doc, [
        "sa région",
        "sa culture",
        "sa surface, en hectares",
        "son niveau d'engrais",
        "son mode d'irrigation",
    ])
    par(doc, "Il obtient deux chiffres : le rendement attendu, avec une fourchette, "
             "et la récolte totale sur sa parcelle, exprimée en sacs de 50 kg.")
    par(doc, "Ce que l'application ne fait pas, volontairement :", gras=True)
    puces(doc, [
        "Elle ne prédit pas une année précise. Elle donne un rendement en conditions normales.",
        "Elle ne demande pas de données techniques comme le NDVI ou le pH.",
        "Elle ne remplace pas un conseiller agricole.",
    ])

    # ---- 4. Les données ----
    titre(doc, "4. Les données")
    par(doc, "Les sources", gras=True)
    tableau(doc,
            ["Source", "Ce qu'elle apporte"],
            [
                ["FAOSTAT", "Rendements nationaux réels du Sénégal, par culture, 2000 à 2022"],
                ["NASA POWER", "Météo réelle par région, saison des pluies de juin à octobre"],
                ["DAPSA", "Facteurs de répartition régionale des rendements"],
            ])
    par(doc, "Un point important, écrit clairement", gras=True)
    par(doc, "Les rendements par parcelle ne sont pas des mesures de terrain. Ils "
             "sont obtenus en répartissant les rendements nationaux FAOSTAT entre les "
             "régions, puis en ajoutant une variabilité liée aux pratiques. Le modèle "
             "apprend donc cette règle de répartition. Ses métriques mesurent sa "
             "capacité à la retrouver, pas à prédire un rendement réellement observé.")
    par(doc, "Ce choix est assumé : il n'existe pas de jeu de données public de "
             "rendements parcellaires au Sénégal. La météo, elle, est bien réelle.")

    # ---- 5. Le modèle ----
    titre(doc, "5. Le modèle")
    par(doc, "L'algorithme retenu est XGBoost, une méthode d'arbres de décision "
             "renforcés, adaptée aux tableaux de données comme le nôtre.")
    par(doc, "Deux précautions de méthode", gras=True)
    puces(doc, [
        "Le découpage entre entraînement et test est fait par groupe année, région "
        "et culture. Des parcelles issues d'une même situation ne se retrouvent pas "
        "des deux côtés, ce qui donnerait des scores trop flatteurs.",
        "Le nombre d'arbres est réglé sur un jeu de validation distinct du jeu de "
        "test. Le test reste ainsi vraiment indépendant.",
    ])
    par(doc, "Résultats mesurés", gras=True)
    tableau(doc,
            ["Indicateur", "Valeur"],
            [
                ["Coefficient R²", metriques.get("r2")],
                ["Erreur RMSE (t/ha)", metriques.get("rmse")],
                ["Erreur moyenne MAE (t/ha)", metriques.get("mae")],
                ["Observations", metriques.get("n_samples")],
                ["Dont entraînement / test", f"{metriques.get('n_train')} / {metriques.get('n_test')}"],
                ["Nombre de variables", metriques.get("n_features")],
            ])
    par(doc, "Ces chiffres sont lus directement dans le modèle entraîné au moment de "
             "la génération de ce rapport ; ils ne sont pas recopiés à la main.",
        italique=True)

    # ---- 6. Architecture ----
    titre(doc, "6. Architecture technique")
    tableau(doc,
            ["Élément", "Rôle"],
            [
                ["streamlit_app.py", "Le simulateur, sur un seul écran"],
                ["pages/1_Le_modele.py", "Traçabilité, métriques et limites du modèle"],
                ["models/train_model.py", "Entraînement et découpage par groupe"],
                ["models/predict.py", "Simulation et validation des entrées"],
                ["utils/referentiel.py", "Source unique des valeurs autorisées"],
                ["data/pipelines/", "Récupération NASA POWER et construction du jeu de données"],
                ["tests/", "Critères d'acceptation automatisés"],
            ])

    # ---- 7. Qualité ----
    titre(doc, "7. Qualité et industrialisation")
    par(doc, "Chaque exigence du cahier des charges correspond à un test "
             "automatique. La chaîne d'intégration continue échoue si l'un d'eux "
             "échoue.")
    tableau(doc,
            ["Critère", "Ce qu'il garantit"],
            [
                ["CA-1", "Toute valeur de l'interface est connue du modèle, sinon erreur explicite"],
                ["CA-2", "Chaque entrée exposée a un effet mesurable et dans le bon sens"],
                ["CA-3", "Le modèle expose sa version, son commit et l'empreinte de ses données"],
                ["CA-4", "Aucune métrique n'est écrite en dur dans le code"],
                ["CA-5", "Les rendements restent dans les plages agronomiques connues"],
                ["CA-6", "La chaîne d'intégration échoue si un critère échoue"],
            ])
    par(doc, f"Le modèle est traçable : version {infos.get('version')}, commit "
             f"{infos.get('commit')}, empreinte de données "
             f"{(infos.get('empreinte_donnees') or '')[:16]}. Ces trois éléments "
             f"suffisent à le reconstruire à l'identique.")
    par(doc, "L'application est aussi disponible en image Docker, dont la CI vérifie "
             "qu'elle démarre correctement.")

    # ---- 8. Limites ----
    titre(doc, "8. Limites reconnues")
    puces(doc, [
        "Les rendements sont répartis à partir de statistiques nationales, pas observés.",
        "Le modèle ne prédit pas une année future.",
        "Les effets climatiques ne sont pas modélisés ; c'est la principale piste d'évolution.",
        "Sept régions et cinq cultures seulement.",
        "Le simulateur donne un ordre de grandeur, il ne remplace pas un conseiller.",
    ])

    # ---- 9. Conclusion ----
    titre(doc, "9. Conclusion")
    par(doc, "AgroPredict SN est un outil simple et utile : il transforme des "
             "statistiques agricoles en une estimation lisible pour un agriculteur. "
             "Sa valeur ne tient pas à un score de précision élevé, mais à sa "
             "rigueur : des données dont l'origine est assumée, un modèle traçable, "
             "et des tests qui garantissent réellement son comportement.")

    doc.save(SORTIE)
    print(f"Rapport généré : {SORTIE}")


if __name__ == "__main__":
    construire()
